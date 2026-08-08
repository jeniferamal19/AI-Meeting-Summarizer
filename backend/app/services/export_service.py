import io

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem

from app.models.schemas import MeetingSummary


def build_pdf(summary: MeetingSummary) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("Meeting Notes", styles["Title"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Overview", styles["Heading2"]))
    story.append(Paragraph(summary.summary, styles["BodyText"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Key Decisions", styles["Heading2"]))
    if summary.key_decisions:
        story.append(ListFlowable(
            [ListItem(Paragraph(d, styles["BodyText"])) for d in summary.key_decisions]
        ))
    else:
        story.append(Paragraph("No decisions recorded.", styles["BodyText"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Risks & Blockers", styles["Heading2"]))
    if summary.risks_blockers:
        story.append(ListFlowable(
            [ListItem(Paragraph(r, styles["BodyText"])) for r in summary.risks_blockers]
        ))
    else:
        story.append(Paragraph("None flagged.", styles["BodyText"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Action Items", styles["Heading2"]))
    if summary.action_items:
        rows = []
        for item in summary.action_items:
            owner = item.owner or "Unassigned"
            due = item.due or "No due date"
            rows.append(ListItem(Paragraph(f"{item.task} — {owner} ({due})", styles["BodyText"])))
        story.append(ListFlowable(rows))
    else:
        story.append(Paragraph("No action items recorded.", styles["BodyText"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Topics Discussed", styles["Heading2"]))
    story.append(Paragraph(", ".join(summary.topics_discussed) or "None.", styles["BodyText"]))

    doc.build(story)
    buffer.seek(0)
    return buffer.read()


def build_docx(summary: MeetingSummary) -> bytes:
    buffer = io.BytesIO()
    doc = Document()

    doc.add_heading("Meeting Notes", level=0)

    doc.add_heading("Overview", level=1)
    doc.add_paragraph(summary.summary)

    doc.add_heading("Key Decisions", level=1)
    if summary.key_decisions:
        for d in summary.key_decisions:
            doc.add_paragraph(d, style="List Bullet")
    else:
        doc.add_paragraph("No decisions recorded.")

    doc.add_heading("Risks & Blockers", level=1)
    if summary.risks_blockers:
        for r in summary.risks_blockers:
            doc.add_paragraph(r, style="List Bullet")
    else:
        doc.add_paragraph("None flagged.")

    doc.add_heading("Action Items", level=1)
    if summary.action_items:
        for item in summary.action_items:
            owner = item.owner or "Unassigned"
            due = item.due or "No due date"
            doc.add_paragraph(f"{item.task} — {owner} ({due})", style="List Bullet")
    else:
        doc.add_paragraph("No action items recorded.")

    doc.add_heading("Topics Discussed", level=1)
    doc.add_paragraph(", ".join(summary.topics_discussed) or "None.")

    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()